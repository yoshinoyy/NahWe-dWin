"""Generate the shareable Steady hackathon concept document."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "docs" / "Steady-Hackathon-Concept.docx"


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def build_document():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.color.rgb = RGBColor(24, 93, 95)
    styles["Heading 1"].font.color.rgb = RGBColor(24, 93, 95)
    styles["Heading 2"].font.color.rgb = RGBColor(57, 117, 119)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Steady")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Student Stress & Workload Manager")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(57, 117, 119)
    tagline = document.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tagline.add_run("Turn overwhelm into one manageable next step.").italic = True
    document.add_paragraph(
        "Hackathon concept brief | Student-focused responsive web MVP | Flask + vanilla JavaScript"
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(
        "Steady is a calm, student-friendly web app that helps people move from “I have too much to do” to one realistic next action. It combines a short self-check-in with assignments, deadlines, estimated effort, and available time. Steady then recommends a capacity-aware focus task and an optional short reset."
    )
    document.add_paragraph(
        "The product is intentionally non-clinical: it does not diagnose stress or make claims about mental health. Its job is to reduce decision overload and make starting easier."
    )

    document.add_heading("The Problem", level=1)
    document.add_paragraph(
        "Students commonly manage competing assignments, exams, clubs, work, and personal responsibilities through scattered notes and large to-do lists. When energy is low or stress is high, a list that says “do everything” can increase avoidance instead of helping someone start."
    )

    document.add_heading("The Solution", level=1)
    document.add_paragraph(
        "Steady asks three practical questions: How are you feeling? How much time and energy do you have? What work is due? It converts those answers into a small plan that respects capacity, highlights urgency, and gives the user permission to adjust."
    )

    document.add_heading("Core MVP", level=1)
    add_bullets(document, [
        "Quick check-in for stress, energy, and available time.",
        "Task inbox with title, deadline, estimated minutes, and importance.",
        "Workload snapshot comparing open effort with today's capacity.",
        "Today's focus card with the next recommended task and a 25-minute starter session.",
        "Reset toolkit with optional breathing, water, stretch, and screen-break prompts.",
        "Local browser persistence without accounts, a database, or external API keys.",
    ])

    document.add_heading("How Prioritization Works", level=1)
    add_bullets(document, [
        "Urgency increases as a deadline gets closer.",
        "Importance gives high-impact assignments additional weight.",
        "Short tasks receive a small boost when stress is high or energy is low, making the first step achievable.",
        "The plan is capped by the user's available time; it does not promise to finish an impossible workload.",
        "High stress produces a smaller recommendation and surfaces a short reset instead of adding pressure.",
    ])

    document.add_heading("User Flow", level=1)
    flow = document.add_table(rows=1, cols=3)
    flow.alignment = WD_TABLE_ALIGNMENT.CENTER
    flow.style = "Light Shading Accent 1"
    for cell, text in zip(flow.rows[0].cells, ["Step", "User action", "Product response"]):
        set_cell_text(cell, text, bold=True)
    for row in [
        ("1", "Completes a 10-second check-in", "Shows a supportive status and capacity summary"),
        ("2", "Adds assignments and deadlines", "Calculates effort, urgency, and workload status"),
        ("3", "Opens Today's Focus", "Recommends one next task and a 25-minute starter"),
        ("4", "Starts or completes the task", "Updates progress and celebrates forward motion"),
        ("5", "Chooses a reset, if needed", "Provides a low-friction recovery prompt"),
    ]:
        cells = flow.add_row().cells
        for cell, text in zip(cells, row):
            set_cell_text(cell, text)

    document.add_heading("Technology Plan", level=1)
    add_bullets(document, [
        "Flask serves the single-page application and a lightweight health endpoint.",
        "HTML and CSS provide the responsive dashboard and calm visual system.",
        "Vanilla JavaScript handles state, form validation, planning logic, and localStorage persistence.",
        "python-docx generates this Word brief from a reproducible script.",
        "No personal data leaves the browser in the MVP.",
    ])

    document.add_heading("8-Hour Hackathon Schedule", level=1)
    schedule = document.add_table(rows=1, cols=2)
    schedule.alignment = WD_TABLE_ALIGNMENT.CENTER
    schedule.style = "Light Shading Accent 1"
    for cell, text in zip(schedule.rows[0].cells, ["Time", "Outcome"]):
        set_cell_text(cell, text, bold=True)
    for row in [
        ("Hour 0–1", "Dashboard shell and data model"),
        ("Hour 1–3", "Tasks, deadlines, persistence"),
        ("Hour 3–5", "Check-in and capacity-aware recommendation"),
        ("Hour 5–6", "Reset toolkit and responsive polish"),
        ("Hour 6–7", "Sample scenarios, accessibility, edge cases"),
        ("Hour 7–8", "Demo rehearsal, screenshots, DOCX export"),
    ]:
        cells = schedule.add_row().cells
        for cell, text in zip(cells, row):
            set_cell_text(cell, text)

    document.add_heading("Demo Script", level=1)
    document.add_paragraph(
        "“Meet Steady. A student has three assignments, two hours, and high stress. Instead of throwing the whole list back at them, Steady checks capacity, identifies the most urgent useful task, and turns it into a 25-minute start. The student can complete it, see the workload update, and take a short reset before continuing. Steady does not diagnose anyone—it makes the next step clearer.”"
    )

    document.add_heading("Future Possibilities", level=1)
    add_bullets(document, [
        "Weekly workload trends and lighter-day suggestions.",
        "Calendar import and export.",
        "Optional AI reflection that rewrites overwhelming tasks into smaller next actions, with an offline fallback.",
        "Campus support links selected by region, with careful privacy and safety review.",
    ])

    document.add_heading("Safety and Scope", level=1)
    document.add_paragraph(
        "Steady is a productivity prototype, not a medical or mental-health diagnostic tool. A production version should be reviewed with student-support professionals, provide region-appropriate support resources, and make privacy choices explicit."
    )

    footer = document.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Built for a hackathon prototype • Steady").italic = True
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_document()
